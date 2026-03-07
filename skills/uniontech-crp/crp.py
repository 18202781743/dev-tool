#!/usr/bin/env python3
"""
uniontech-crp Skill - CRP 打包操作

直接调用 CRP HTTP API，不依赖任何外部二进制工具。

用法：
  python crp.py pack     --topic yss-dde-test --name deepin-terminal
  python crp.py test     --topic yss-dde-test --name deepin-terminal
  python crp.py topics   --topic yss
  python crp.py projects --name deepin-terminal
  python crp.py instances --topic yss-dde-test
  python crp.py branches --topic yss-dde-test --name deepin-terminal
  python crp.py gendoc   --topic yss-dde-test [--template path/to/tpl.xlsx]

Exit Codes:
  0 - 成功
  1 - 一般错误
  2 - 参数错误
  3 - 认证错误
"""

from __future__ import annotations

import json
import re
import sys
import argparse
import os
from typing import Any, Dict, List, Optional

import requests

from auth import get_cached_credentials, login_with_credentials, login_with_cached_credentials
from config import get_config
from security import InputValidator, ValidationResult

# Exit codes
EXIT_SUCCESS = 0
EXIT_ERROR   = 1
EXIT_USAGE   = 2
EXIT_AUTH    = 3

CRP_BASE_URL  = "https://crp.uniontech.com"
HEADERS_JSON  = {"Content-Type": "application/json"}
SHUTTLE_BASE  = "https://shuttle.uniontech.com/#/tasks/task?taskid="


def error_response(error_type: str, message: str, suggestion: str = None, **extra) -> str:
    resp = {"error": error_type, "message": message}
    if suggestion:
        resp["suggestion"] = suggestion
    resp.update(extra)
    return json.dumps(resp, ensure_ascii=False, indent=2)


# ============================================================================
# CRP API 低层封装
# ============================================================================

class _CRPAPI:
    """CRP HTTP API 封装（内部使用，token 由外部注入）"""

    def __init__(self, token: str, cfg: Dict[str, Any]):
        self.token = token
        self.cfg   = cfg

    def _auth_headers(self) -> Dict[str, str]:
        return {**HEADERS_JSON, "Authorization": f"Bearer {self.token}"}

    def fetch_user(self) -> str:
        resp = requests.get(
            f"{CRP_BASE_URL}/api/user",
            headers=self._auth_headers(), timeout=30,
        )
        resp.raise_for_status()
        return resp.json().get("Name", "")

    def list_topics(self, keyword: str) -> List[Dict]:
        data = {
            "TopicType": self.cfg["topic_type"],
            "UserName":  self.fetch_user(),
            "BranchID":  self.cfg["branch_id"],
        }
        resp = requests.post(
            f"{CRP_BASE_URL}/api/topics/search",
            headers=self._auth_headers(), json=data, timeout=30,
        )
        resp.raise_for_status()
        return [
            {"id": t["ID"], "name": t["Name"]}
            for t in resp.json()
            if re.search(keyword, t.get("Name", ""), re.IGNORECASE)
        ]

    def list_projects(self, name_filter: str) -> List[Dict]:
        data = {
            "page": 0, "perPage": 0,
            "projectGroupID": 0, "newCommit": False, "archived": False,
            "branchID": self.cfg["branch_id"],
            "name": name_filter,
        }
        resp = requests.post(
            f"{CRP_BASE_URL}/api/project",
            headers=self._auth_headers(), json=data, timeout=30,
        )
        resp.raise_for_status()
        return [
            {"id": p["ID"], "name": p["Name"], "url": p.get("RepoUrl", "")}
            for p in resp.json().get("Projects", [])
            if re.search(name_filter or ".", p.get("Name", ""), re.IGNORECASE)
        ]

    def _fetch_commit_message(self, repo_url: str, commit: str) -> str:
        try:
            resp = requests.post(
                f"{CRP_BASE_URL}/api/projects/getGerritCommitMessage",
                headers=self._auth_headers(),
                json={"repo_url": repo_url, "commit_id": commit},
                timeout=30,
            )
            resp.raise_for_status()
            return resp.json().get("message", "")
        except Exception:
            return ""

    def list_branches(self, project_id: int, project_url: str,
                      branch_filter: str) -> List[Dict]:
        resp = requests.get(
            f"{CRP_BASE_URL}/api/projects/{project_id}/branches",
            headers=self._auth_headers(), timeout=30,
        )
        resp.raise_for_status()
        result = []
        for b in resp.json():
            name   = b.get("Name", "")
            commit = b.get("Commit", "")
            if branch_filter and not re.search(branch_filter, name, re.IGNORECASE):
                continue
            msg = self._fetch_commit_message(project_url, commit) or b.get("Message", "")
            result.append({"name": name, "commit": commit, "changelog": msg})
        return result

    def list_created_instances(self, topic_id: int) -> List[Dict]:
        resp = requests.get(
            f"{CRP_BASE_URL}/api/topics/{topic_id}/releases",
            headers=self._auth_headers(), timeout=30,
        )
        resp.raise_for_status()
        result = []
        for i in resp.json():
            build_id    = i.get("BuildID", 0)
            build_state = i.get("BuildState", {}).get("state", "UNKNOWN")
            # UPLOAD_OK 时构建成功，无需日志链接；其他状态给 shuttle 链接方便排查
            build_url   = "" if build_state == "UPLOAD_OK" else (
                f"{SHUTTLE_BASE}{build_id}" if build_id else ""
            )
            result.append({
                "id":           i["ID"],
                "project_id":   i["ProjectID"],
                "project_name": i["ProjectName"],
                "branch":       i["Branch"],
                "tag":          i["Tag"],
                "build_state":  build_state,
                "build_id":     build_id,
                "build_url":    build_url,
            })
        return result

    def delete_instance(self, instance_id: int) -> None:
        resp = requests.delete(
            f"{CRP_BASE_URL}/api/topic_releases/{instance_id}",
            headers=self._auth_headers(), timeout=30,
        )
        resp.raise_for_status()

    def create_instance(self, topic_id: int, payload: Dict) -> Dict:
        resp = requests.post(
            f"{CRP_BASE_URL}/api/topics/{topic_id}/new_release",
            headers=self._auth_headers(), json=payload, timeout=30,
        )
        resp.raise_for_status()
        return resp.json() if resp.text.strip() else {}

    def get_topic_urls(self, topic_id: int) -> List[str]:
        resp = requests.post(
            f"{CRP_BASE_URL}/api/topic_urls",
            headers=self._auth_headers(),
            json={"branchID": self.cfg["branch_id"], "topicID": topic_id},
            timeout=30,
        )
        resp.raise_for_status()
        data = resp.json()
        urls = data.get("RepoUrls", [])
        if isinstance(urls, str):
            return [urls] if urls else []
        return urls or []


# ============================================================================
# CRPClient - 对外接口
# ============================================================================

class CRPClient:
    """
    CRP 打包操作客户端。

    token 由调用方（login_with_cached_credentials 的返回值）注入，
    CRPClient 本身不感知认证细节。
    """

    def __init__(self, token: str, cfg: Optional[Dict[str, Any]] = None):
        self._api = _CRPAPI(token, cfg or get_config())

    # ------------------------------------------------------------------
    # 对外操作
    # ------------------------------------------------------------------

    def pack(
        self,
        topic: str,
        name: str,
        branch: Optional[str] = None,
        tag: Optional[str] = None,
    ) -> Dict[str, Any]:
        """
        CRP 打包（创建或替换 instance）。

        Args:
            topic:  主题名（如 yss-dde-test）
            name:   包名（如 deepin-terminal）
            branch: 分支（不填则用配置中的 default_branch）
            tag:    指定 tag（不填则使用 changelog 模式自动递增）
        """
        for val, fn in [
            (topic, InputValidator.validate_topic),
            (name,  InputValidator.validate_package_name),
        ]:
            r: ValidationResult = fn(val)
            if not r.valid:
                return {"success": False, "msg": r.message}

        cfg     = self._api.cfg
        _branch = branch or cfg["default_branch"]
        r = InputValidator.validate_branch(_branch)
        if not r.valid:
            return {"success": False, "msg": r.message}

        api      = self._api
        topics   = api.list_topics(topic)
        projects = api.list_projects(name)

        if not topics:
            return {"success": False, "msg": f"未找到主题: {topic!r}"}
        if not projects:
            return {"success": False, "msg": f"未找到项目: {name!r}"}

        created = []
        errors  = []
        for t in topics:
            for p in projects:
                branches = api.list_branches(p["id"], p["url"], _branch)
                if not branches:
                    errors.append(f"{p['name']}: 未找到分支 {_branch!r}")
                    continue
                for b in branches:
                    # 替换已存在的 instance
                    for inst in api.list_created_instances(t["id"]):
                        if inst["project_name"] == p["name"] and inst["branch"] == b["name"]:
                            api.delete_instance(inst["id"])

                    payload = {
                        "Arches":         cfg["archs"],
                        "BaseTag":        None,
                        "Branch":         b["name"],
                        "BuildID":        0,
                        "BuildState":     None,
                        "Changelog":      [b["changelog"]],
                        "Commit":         b["commit"],
                        "History":        None,
                        "ID":             0,
                        "ProjectID":      p["id"],
                        "ProjectName":    p["name"],
                        "ProjectRepoUrl": p["url"],
                        "SlaveNode":      None,
                        "Tag":            tag or "",
                        "TagSuffix":      None,
                        "TopicID":        t["id"],
                        "TopicType":      cfg["topic_type"],
                        "ChangeLogMode":  tag is None,
                        "RepoType":       "deb",
                        "Custom":         True,
                        "BranchID":       str(cfg["branch_id"]),
                    }
                    api.create_instance(t["id"], payload)
                    created.append(f"{t['name']} / {p['name']} / {b['name']}")

        if created:
            return {"success": True, "msg": f"已提交打包: {', '.join(created)}"}
        return {
            "success": False,
            "msg":     "\n".join(errors) or "未匹配到任何 topic/project/branch 组合",
        }

    def test(
        self,
        topic: str,
        name: str,
        branch: Optional[str] = None,
    ) -> Dict[str, Any]:
        """测试模式：列出将要打包的组合，不实际提交"""
        for val, fn in [
            (topic, InputValidator.validate_topic),
            (name,  InputValidator.validate_package_name),
        ]:
            r: ValidationResult = fn(val)
            if not r.valid:
                return {"success": False, "msg": r.message}

        _branch = branch or self._api.cfg["default_branch"]
        api     = self._api
        items   = []

        for t in api.list_topics(topic):
            for p in api.list_projects(name):
                for b in api.list_branches(p["id"], p["url"], _branch):
                    items.append({
                        "topic":     t["name"],
                        "project":   p["name"],
                        "branch":    b["name"],
                        "changelog": b["changelog"],
                    })

        if items:
            return {"success": True, "msg": f"将打包 {len(items)} 个组合", "items": items}
        return {"success": False, "msg": "未找到匹配的 topic/project/branch 组合"}

    def topics(self, keyword: str) -> Dict[str, Any]:
        r = InputValidator.validate_topic(keyword)
        if not r.valid:
            return {"success": False, "msg": r.message}
        result = self._api.list_topics(keyword)
        return {"success": True, "topics": result, "msg": f"找到 {len(result)} 个主题"}

    def projects(self, name: Optional[str] = None) -> Dict[str, Any]:
        keyword = name or ""
        if keyword:
            r = InputValidator.validate_package_name(keyword)
            if not r.valid:
                return {"success": False, "msg": r.message}
        result = self._api.list_projects(keyword)
        return {"success": True, "projects": result, "msg": f"找到 {len(result)} 个项目"}

    def instances(self, topic: str) -> Dict[str, Any]:
        r = InputValidator.validate_topic(topic)
        if not r.valid:
            return {"success": False, "msg": r.message}
        api    = self._api
        topics = api.list_topics(topic)
        if not topics:
            return {"success": False, "msg": f"未找到主题: {topic!r}"}
        all_instances = []
        for t in topics:
            for inst in api.list_created_instances(t["id"]):
                inst["topic_name"] = t["name"]
                all_instances.append(inst)
        return {
            "success": True,
            "instances": all_instances,
            "msg": f"共 {len(all_instances)} 个 instance",
        }

    def branches(self, topic: str, name: str) -> Dict[str, Any]:
        for val, fn in [(topic, InputValidator.validate_topic),
                        (name,  InputValidator.validate_package_name)]:
            r = fn(val)
            if not r.valid:
                return {"success": False, "msg": r.message}
        api      = self._api
        topics   = api.list_topics(topic)
        projects = api.list_projects(name)
        if not topics:
            return {"success": False, "msg": f"未找到主题: {topic!r}"}
        if not projects:
            return {"success": False, "msg": f"未找到项目: {name!r}"}
        all_branches = []
        for p in projects:
            for b in api.list_branches(p["id"], p["url"], ""):
                b["project"] = p["name"]
                all_branches.append(b)
        return {
            "success": True,
            "branches": all_branches,
            "msg": f"共 {len(all_branches)} 个分支",
        }

    def gendoc(self, topic: str, template: Optional[str] = None) -> Dict[str, Any]:
        r = InputValidator.validate_topic(topic)
        if not r.valid:
            return {"success": False, "msg": r.message}
        if template:
            r = InputValidator.validate_file_path(template)
            if not r.valid:
                return {"success": False, "msg": r.message}
        return self._gendoc(topic, template)

    def _gendoc(self, topic_keyword: str, template_path: Optional[str]) -> Dict[str, Any]:
        try:
            from openpyxl import load_workbook
        except ImportError:
            return {"success": False, "msg": "缺少依赖: openpyxl，请运行 pip install openpyxl"}

        from pathlib import Path as _Path

        if template_path and os.path.exists(template_path):
            tpl = template_path
        else:
            candidates = [
                os.path.join(os.getcwd(), "crp-gendoc.xlsx"),
                str(_Path.home() / ".config" / "uniontech-oa" / "crp-gendoc.xlsx"),
                str(_Path.home() / ".openclaw" / "config" / "uniontech-oa" / "crp-gendoc.xlsx"),
            ]
            tpl = next((p for p in candidates if os.path.exists(p)), None)
            if not tpl:
                return {"success": False, "msg": "未找到模板文件 crp-gendoc.xlsx，请放置到配置目录或用 --template 指定"}

        api    = self._api
        topics = api.list_topics(topic_keyword)
        if not topics:
            return {"success": False, "msg": f"未找到主题: {topic_keyword!r}"}
        topic = topics[0]

        releases_resp = requests.get(
            f"{CRP_BASE_URL}/api/topics/{topic['id']}/releases",
            headers=api._auth_headers(), timeout=30,
        )
        releases_resp.raise_for_status()
        releases = releases_resp.json()
        if not releases:
            return {"success": True, "msg": "无 release 数据，未生成文档"}

        wb = load_workbook(tpl)
        if "ChangeLog" not in wb.sheetnames:
            return {"success": False, "msg": "模板中缺少 ChangeLog 工作表"}
        ws = wb["ChangeLog"]

        module_name_row = test_desc_row = pack_branch_row = None
        for row_idx in range(1, 100):
            val = ws.cell(row=row_idx, column=1).value
            if val and "模块名" in str(val):
                module_name_row = row_idx
            elif val == "转测说明":
                test_desc_row = row_idx
            elif val and "打包分支及测试主题" in str(val):
                pack_branch_row = row_idx

        if not module_name_row or not test_desc_row:
            return {"success": False, "msg": "模板格式异常：未找到关键行"}

        insert_row    = module_name_row + 1
        reserved_rows = test_desc_row - insert_row
        if len(releases) > reserved_rows:
            return {"success": False,
                    "msg": f"数据行数({len(releases)})超过模板预留行数({reserved_rows})"}

        for idx, release in enumerate(releases):
            row = insert_row + idx
            ws.cell(row=row, column=1, value=release.get("SourcePkgName", ""))
            ws.cell(row=row, column=2, value=release.get("SourcePkgName", ""))
            ws.cell(row=row, column=3, value=release.get("Tag", ""))
            ws.cell(row=row, column=4, value=release.get("Commit", ""))

        if pack_branch_row:
            repo_urls = api.get_topic_urls(topic["id"])
            if repo_urls:
                ws.cell(row=pack_branch_row, column=2, value=", ".join(repo_urls))

        filename = f"测试-桌面专业版-转测申请单-{topic['name']}.xlsx"
        wb.save(filename)
        return {"success": True, "msg": f"文档已生成: {filename}", "file": filename}


# ============================================================================
# CLI 入口
# ============================================================================

def main() -> int:
    parser = argparse.ArgumentParser(description="CRP 打包操作（直接调用 CRP API）")
    parser.add_argument("action", choices=[
        "pack", "test", "topics", "projects", "instances", "branches", "gendoc"
    ])
    parser.add_argument("--topic",    help="测试主题（如 yss-dde-test）")
    parser.add_argument("--name",     help="包名")
    parser.add_argument("--branch",   help="分支（默认读取配置）")
    parser.add_argument("--tag",      help="指定 tag（不填则自动递增）")
    parser.add_argument("--template", help="转测文档模板路径")
    args = parser.parse_args()

    # 认证：用缓存的 OA 凭证登录 CRP 获取 token
    cred = login_with_cached_credentials()
    if not cred:
        print(error_response(
            "auth_error", "未找到 OA 凭证缓存",
            "请提供 OA 账号和密码，然后调用 login_with_credentials(loginid, password)"
        ), file=sys.stderr)
        return EXIT_AUTH

    client = CRPClient(cred["token"])

    try:
        if args.action == "pack":
            if not args.topic or not args.name:
                print(error_response("usage_error", "--topic 和 --name 均必填"), file=sys.stderr)
                return EXIT_USAGE
            result = client.pack(args.topic, args.name, args.branch, args.tag)

        elif args.action == "test":
            if not args.topic or not args.name:
                print(error_response("usage_error", "--topic 和 --name 均必填"), file=sys.stderr)
                return EXIT_USAGE
            result = client.test(args.topic, args.name, args.branch)

        elif args.action == "topics":
            if not args.topic:
                print(error_response("usage_error", "--topic 必填"), file=sys.stderr)
                return EXIT_USAGE
            result = client.topics(args.topic)

        elif args.action == "projects":
            result = client.projects(args.name)

        elif args.action == "instances":
            if not args.topic:
                print(error_response("usage_error", "--topic 必填"), file=sys.stderr)
                return EXIT_USAGE
            result = client.instances(args.topic)

        elif args.action == "branches":
            if not args.topic or not args.name:
                print(error_response("usage_error", "--topic 和 --name 均必填"), file=sys.stderr)
                return EXIT_USAGE
            result = client.branches(args.topic, args.name)

        elif args.action == "gendoc":
            if not args.topic:
                print(error_response("usage_error", "--topic 必填"), file=sys.stderr)
                return EXIT_USAGE
            result = client.gendoc(args.topic, args.template)

    except Exception as e:
        print(error_response("api_error", str(e)), file=sys.stderr)
        return EXIT_ERROR

    print(json.dumps(result, ensure_ascii=False, indent=2))
    return EXIT_SUCCESS if result.get("success") else EXIT_ERROR


if __name__ == "__main__":
    for var in ["https_proxy", "http_proxy", "all_proxy", "HTTPS_PROXY", "HTTP_PROXY"]:
        os.environ.pop(var, None)
    sys.exit(main())
